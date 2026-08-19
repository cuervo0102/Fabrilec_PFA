import docx


def extract_docx(path: str) -> dict:
    doc = docx.Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return {"text": "\n".join(parts), "method": "docx", "n_pages": None}